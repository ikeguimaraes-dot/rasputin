"""Exporta exclusivamente o snapshot salvo, sem recalcular com regras novas."""

import html
import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from auditoria.reports import safe_cell


def export(row, kind):
    result = row["result_payload"]
    source = row["input_snapshot"]
    fields = row["catalog_snapshot"]["modules"][row["module"]]["fields"]
    labels = {f["key"]: f["label"] for f in fields}
    headers = [
        "Tributo",
        "Base R$",
        "Débito R$",
        "Deduções R$",
        "Saldo a recolher R$",
        "Excedente R$",
    ]
    values = [
        [
            (x[k] if x[k] is not None else "—")
            for k in ("tax", "base", "gross", "deductions", "due", "remaining")
        ]
        for x in result["lines"]
    ]
    context = [
        ["Cliente", source["client"]["razao_social"]],
        ["CNPJ", source["client"]["cnpj"]],
        ["Competência", str(row["period"])],
        ["Periodicidade", result["period_type"]],
        ["Versão", result["catalog_version"]],
        ["Hash da apuração", row["sha256"]],
        ["Hash das regras", result["catalog_hash"]],
        ["Escopo", result["scope"]],
        ["Situação", result["status"]],
    ]
    inputs = [[labels.get(k, k), str(v)] for k, v in source["values"].items()]
    sources = [
        [s["title"], s["url"], s["sha256"], s.get("capture_kind", "")] for s in result["sources"]
    ]
    notes = [
        result.get("notice", ""),
        *result["warnings"],
        *(m["label"] for m in result.get("missing", [])),
    ]
    evidence = result.get("memory", {}).get("documents", {})
    document_rows = [
        [
            "Nota",
            "Item",
            "CFOP",
            "Descrição",
            "Operação",
            "ICMS",
            "Base esperada",
            "Base PIS",
            "Base COFINS",
            "Situação",
        ]
    ] + [
        [
            r.get(k, "")
            for k in (
                "document",
                "item",
                "cfop",
                "description",
                "operation",
                "icms",
                "expected_base",
                "pis_base",
                "cofins_base",
                "status",
            )
        ]
        for r in evidence.get("rows", [])
    ]
    subtotals = [
        ["Tributo", "Base recomposta", "Débito às alíquotas declaradas", "Itens", "Sem dados"]
    ] + [
        [
            r["tax"],
            r["base"] if r["items"] else "Não calculado",
            r["gross"] if r["items"] else "Não calculado",
            r["items"],
            r["missing"],
        ]
        for r in evidence.get("recalculated", [])
    ]
    if evidence:
        context.extend(
            [
                [
                    "Conferência de bases",
                    f"{evidence['base_review']['missing_icms']} linhas sem excluir ICMS; "
                    f"{evidence['base_review']['other_difference']} com outra divergência; "
                    f"{evidence['base_review']['compatible']} compatíveis; "
                    f"{evidence['base_review']['unassessed']} sem dados suficientes.",
                ],
                ["Hash documental", evidence["hash"]],
            ]
        )
        notes.extend(f"Origem: {u['name']} · SHA-256 {u['sha256']}" for u in evidence["uploads"])
        notes.append(
            "Débitos documentais usam as alíquotas declaradas, antes de créditos e retenções. "
            "Não certificam enquadramento nem saldo a recolher."
        )
    if kind == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in [
            ("Apuração", context + [headers] + values),
            ("Dados declarados", inputs),
            ("Conferência documental", document_rows),
            ("Débitos documentais", subtotals),
            ("Fontes", [["Norma", "URL", "Hash da captura", "Tipo da captura"]] + sources),
            (
                "Memória",
                [
                    [k, json.dumps(v, ensure_ascii=False)]
                    for k, v in result["memory"].items()
                    if k != "documents"
                ],
            ),
            ("Observações", [[x] for x in notes]),
        ]:
            ws = wb.create_sheet(name)
            for r in rows:
                ws.append([safe_cell(x) for x in r])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(
                    65, max(20, max(len(str(c.value or "")) for c in col) + 2)
                )
        out = BytesIO()
        wb.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # PDF is the concise client report; Excel retains every assessed item.
    document_rows = [document_rows[0]] + [
        r for r in document_rows[1:] if r[-1] in ("missing_icms", "other_difference")
    ]
    status_labels = {
        "missing_icms": "ICMS não excluído da base",
        "other_difference": "Outra divergência de base",
    }
    for r in document_rows[1:]:
        r[-1] = status_labels.get(r[-1], r[-1])
    from weasyprint import HTML
    from weasyprint.urls import URLFetcher

    def e(x):
        return html.escape(str(x))

    def table(rows, style=""):
        if not rows:
            return ""
        return (
            '<table class="'
            + style
            + '">'
            + "".join("<tr>" + "".join("<td>" + e(x) + "</td>" for x in r) + "</tr>" for r in rows)
            + "</table>"
        )

    body = (
        '<html lang="pt-BR"><meta charset="utf-8"><style>@page{size:A4;margin:18mm}'
        "body{font:10pt "
        "sans-serif;color:#173d35}table{width:100%;table-layout:fixed;"
        "border-collapse:collapse;margin:12px 0;font-size:8pt}"
        ".context td:first-child{width:25%}.context{font-size:9pt}"
        "tr{break-inside:avoid}h2{break-after:avoid}.draft{padding:10px;background:#fff4da}"
        "td{border:1px solid "
        "#ddd;padding:6px;overflow-wrap:anywhere}pre{white-space:pre-wrap;overflow-wrap:anywhere}"
        "h1{font-size:20pt}</style><h1>"
        + e(result["title"])
        + "</h1>"
        + (
            "<p class='draft'><b>PRÉVIA NÃO FECHADA</b> — " + e(result.get("notice", "")) + "</p>"
            if result["status"] == "rascunho"
            else ""
        )
        + table(context, "context")
        + (table([headers] + values) if values else "")
        + ("<h2>Débitos documentais parciais</h2>" + table(subtotals) if evidence else "")
        + ("<h2>Dados declarados</h2>" + table(inputs) if inputs else "")
        + "<h2>Memória de cálculo</h2><pre>"
        + e(
            json.dumps(
                {k: v for k, v in result["memory"].items() if k != "documents"},
                ensure_ascii=False,
                indent=2,
            )
        )
        + "</pre>"
        + ("<h2>Conferência documental</h2>" + table(document_rows) if evidence else "")
        + "<h2>Fontes</h2>"
        + table(sources)
        + "".join("<p>" + e(x) + "</p>" for x in notes)
        + "</html>"
    )
    fetcher = URLFetcher(allowed_protocols={"data"}, allow_redirects=False, fail_on_errors=True)
    return HTML(string=body, url_fetcher=fetcher).write_pdf(), "application/pdf"
