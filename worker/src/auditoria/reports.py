"""Relatórios determinísticos; textos não dependem de IA."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from auditoria.confirmed_report import example, publish
from auditoria.domain import Result

FOOTER = "Relatório de conferência. Não substitui a validação do contador responsável."


def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def xlsx_report(result: Result, context: dict) -> bytes:
    result = publish(result)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    for key, val in context.items():
        ws.append([key, safe_cell(str(val))])
    ws.append(["Motor", result.engine_version])
    ws.append(["Hash das entradas", result.input_hash])
    ws.append(["Hash das regras", result.rules_hash])
    ws.append(["Limitação", result.summary["impact_notice"]])
    ws.append(["Rodapé", FOOTER])
    base = result.summary.get("base_review")
    if base:
        ws.append(["Linhas sem descontar o ICMS da base", base["missing_icms"]])
        ws.append(["Linhas com outra divergência na composição da base", base["other_difference"]])
        ws.append(["Linhas compatíveis com a fórmula", base["compatible"]])
        if base["rows"]:
            ws.append(["Exemplo", safe_cell(example(base["rows"][0]))])
        ws = wb.create_sheet("Bases PIS COFINS")
        ws.append(
            [
                "Nota",
                "Linha",
                "Produto",
                "CFOP",
                "Resultado",
                "Valor da operação R$",
                "ICMS R$",
                "Base esperada R$",
                "Base PIS R$",
                "Base COFINS R$",
                "CST PIS",
                "CST COFINS",
                "Arquivo",
                "Aba",
            ]
        )
        for row in base["rows"]:
            ws.append(
                [
                    safe_cell(row["document"]),
                    row["source"].get("line"),
                    safe_cell(row["description"]),
                    row["cfop"],
                    "ICMS não descontado"
                    if row["kind"] == "missing_icms"
                    else "Outra divergência de composição",
                ]
                + [
                    float(row[k]) if row[k] is not None else None
                    for k in ("operation_value", "icms", "expected_base", "pis_base", "cofins_base")
                ]
                + [
                    row["pis_cst"],
                    row["cofins_cst"],
                    safe_cell(row["source"].get("file")),
                    safe_cell(row["source"].get("sheet")),
                ]
            )
    ws = wb.create_sheet("Ocorrências")
    ws.append(
        [
            "Regra",
            "Severidade",
            "Produto",
            "Descrição",
            "Nota",
            "Data",
            "Item",
            "Arquivo",
            "Aba",
            "Linha",
            "Problema",
            "Como está",
            "Como deveria estar",
            "Impacto potencial R$",
            "Direção potencial",
            "Base legal",
            "Fonte",
            "ID regra",
            "ID apontamento",
            "NCM",
            "CFOP",
            "CEST",
            "ICMS CST",
            "ICMS CSOSN",
            "ICMS base",
            "ICMS alíquota %",
            "ICMS valor",
            "PIS CST",
            "PIS base",
            "PIS alíquota %",
            "PIS valor",
            "COFINS CST",
            "COFINS base",
            "COFINS alíquota %",
            "COFINS valor",
        ]
    )
    for f in result.findings:
        if base and f.code == "C09":
            continue
        for ev in f.evidence:
            item, source = ev["item"], ev["item"]["source"]
            row = [
                f.code,
                f.severity,
                f.product,
                f.description,
                ev["document"],
                ev["issued"],
                item["n_item"],
                source.get("file"),
                source.get("sheet"),
                source.get("line"),
                f.message,
                str({k: item.get(k) for k in ("ncm", "cfop", "cest", "icms", "pis", "cofins")}),
                str(f.expected),
                float(f.impact) if f.impact is not None else None,
                f.direction,
                f.legal_basis,
                f.source_url,
                f.rule_id,
                f.id,
                item["ncm"],
                item["cfop"],
                item["cest"],
                item["icms"]["cst"],
                item["icms"]["csosn"],
                *[
                    float(item[t][field]) if item[t][field] is not None else None
                    for t in ("icms",)
                    for field in ("base", "rate", "value")
                ],
                item["pis"]["cst"],
                *[
                    float(item["pis"][field]) if item["pis"][field] is not None else None
                    for field in ("base", "rate", "value")
                ],
                item["cofins"]["cst"],
                *[
                    float(item["cofins"][field]) if item["cofins"][field] is not None else None
                    for field in ("base", "rate", "value")
                ],
            ]
            ws.append([safe_cell(v) for v in row])
    ws = wb.create_sheet("Por CFOP e categoria")
    ws.append(
        [
            "CFOP",
            "Natureza",
            "Categoria declarada",
            "ICMS alíquota %",
            "Itens",
            "Valor dos itens R$",
            "ICMS informado R$",
        ]
    )
    for row in result.summary.get("operations", []):
        ws.append(
            [safe_cell(row.get(k)) for k in ("cfop", "nature", "category")]
            + [
                float(row["icms_rate"]) if row["icms_rate"] is not None else None,
                row["items"],
                float(row["value"]) if row["value"] is not None else None,
                float(row["icms_value"]) if row["icms_value"] is not None else None,
            ]
        )
    ws = wb.create_sheet("Itens conferidos")
    keys = (
        "document",
        "issued",
        "code",
        "description",
        "ncm",
        "cfop",
        "nature",
        "category",
        "icms_cst",
        "pis_cst",
        "cofins_cst",
        "icms_rate",
        "value",
        "icms_value",
    )
    ws.append(
        [
            "Nota",
            "Data",
            "Código",
            "Descrição",
            "NCM",
            "CFOP",
            "Natureza",
            "Categoria declarada",
            "CST ICMS",
            "CST PIS",
            "CST COFINS",
            "ICMS alíquota %",
            "Valor do item R$",
            "ICMS informado R$",
            "Arquivo",
            "Aba",
            "Linha",
        ]
    )
    for row in result.summary.get("review_rows", []):
        ws.append(
            [
                safe_cell(row.get(k))
                if k not in ("icms_rate", "value", "icms_value")
                else (float(row[k]) if row.get(k) is not None else None)
                for k in keys
            ]
            + [safe_cell(row["source"].get(k)) for k in ("file", "sheet", "line")]
        )
    for sheet in wb:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173D35")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(
                65, max(15, max(len(str(c.value or "")) for c in list(column)[:100]) + 2)
            )
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def html_report(result: Result, context: dict, branding: dict | None = None) -> str:
    from auditoria.report_design import client_html

    return client_html(publish(result), context, branding, FOOTER)


def pdf_report(result: Result, context: dict, branding: dict | None = None) -> bytes:
    from weasyprint import HTML
    from weasyprint.urls import URLFetcher

    # Somente imagens PNG/JPEG embutidas e validadas pelo contrato de branding.
    # Bloqueia http(s), file, FTP e redirecionamentos para recursos externos.
    fetcher = URLFetcher(allowed_protocols={"data"}, allow_redirects=False, fail_on_errors=True)
    return HTML(string=html_report(result, context, branding), url_fetcher=fetcher).write_pdf()
