from __future__ import annotations

import csv
import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from zipfile import ZipFile

import openpyxl
import xlrd

from auditoria.domain import Document, Ingestion, Item
from auditoria.ingestion.biff import recover
from auditoria.ingestion.xml import MAX_BYTES, MAX_EXPANDED

FIELDS = {
    "issued": "Data de emissão",
    "key": "Chave da nota",
    "number": "Número da nota",
    "series": "Série",
    "model": "Modelo",
    "n_item": "Número do item",
    "code": "Código do produto",
    "description": "Descrição",
    "ncm": "NCM",
    "cest": "CEST",
    "cfop": "CFOP",
    "unit": "Unidade",
    "quantity": "Quantidade",
    "unit_value": "Valor unitário",
    "value": "Valor do produto",
    "discount": "Desconto",
    "freight": "Frete",
    "insurance": "Seguro",
    "other": "Outras despesas",
    "recipient_taxpayer": "Destinatário contribuinte",
    "recipient_uf": "UF destinatário",
    "final_consumer": "Consumidor final",
    "operation": "Operação (entrada/saida)",
    "status": "Situação (autorizada/cancelada)",
}
for _tax in ("icms", "pis", "cofins", "ipi", "st", "fcp"):
    for _field in ("cst", "csosn", "base", "rate", "value", "reduction"):
        FIELDS[f"{_tax}.{_field}"] = f"{_tax.upper()} — {_field}"


def workbook(data: bytes, filename: str):
    if len(data) > MAX_BYTES:
        raise ValueError("Planilha excede 30 MB")
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=";,\t,")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(StringIO(content), dialect))
        books, warnings, partial = {"CSV": rows}, [], False
    elif ext == "xlsx":
        with ZipFile(BytesIO(data)) as z:
            if sum(m.file_size for m in z.infolist()) > MAX_EXPANDED:
                raise ValueError("XLSX descompactado excede limite")
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        books = {}
        try:
            for sheet in wb:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > 2_000_000:
                    raise ValueError("Planilha excede limite de células")
                books[sheet.title] = list(sheet.values)
        finally:
            wb.close()
        warnings, partial = [], False
    elif ext == "xls":
        try:
            wb = xlrd.open_workbook(file_contents=data)
            books = {}
            for sheet in wb.sheets():
                rows = []
                for r in range(sheet.nrows):
                    row = []
                    for c in range(sheet.ncols):
                        cell = sheet.cell(r, c)
                        row.append(
                            xlrd.xldate_as_datetime(cell.value, wb.datemode)
                            if cell.ctype == xlrd.XL_CELL_DATE
                            else cell.value
                        )
                    rows.append(row)
                books[sheet.name] = rows
            warnings, partial = [], False
        except (xlrd.XLRDError, OSError, IndexError, ValueError, ArithmeticError):
            books, warnings = recover(data)
            partial = True
    else:
        raise ValueError("Formato esperado: XLS, XLSX ou CSV")
    if sum(len(rows) for rows in books.values()) > 100000:
        raise ValueError("Planilha excede 100 mil linhas")
    return books, warnings, partial


def headers(rows, count: int):
    if not 1 <= count <= 10:
        raise ValueError("Cabeçalho deve ter entre 1 e 10 linhas")
    width = max((len(r) for r in rows[:count]), default=0)
    levels = []
    for index, row in enumerate(rows[:count]):
        level, previous = [], ""
        for c in range(width):
            value = str(row[c]).strip() if c < len(row) and row[c] is not None else ""
            if index < count - 1:
                previous = value or previous
                value = previous
            level.append(value)
        levels.append(level)
    names = [
        " / ".join(level[c] for level in levels if level[c]) or f"Coluna {c + 1}"
        for c in range(width)
    ]
    signature = hashlib.sha256(json.dumps(names, ensure_ascii=False).encode()).hexdigest()
    return names, signature


def preview(data, filename, sheet=None, header_rows=1):
    books, warnings, partial = workbook(data, filename)
    sheet = sheet or next(iter(books))
    if sheet not in books:
        raise ValueError("Aba não encontrada")
    names, signature = headers(books[sheet], header_rows)
    return {
        "sheets": list(books),
        "sheet": sheet,
        "headers": names,
        "signature": signature,
        "sample": [
            [str(v) if v is not None else "" for v in r]
            for r in books[sheet][header_rows : header_rows + 5]
        ],
        "rows": max(0, len(books[sheet]) - header_rows),
        "warnings": warnings,
        "partial": partial,
        "fields": FIELDS,
    }


def decimal(value):
    if value is None or str(value).strip() == "":
        return None
    s = str(value).strip().replace("R$", "").replace("%", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    return Decimal(s)


def string(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def boolean(value):
    s = str(value).strip().lower()
    if s in ("1", "true", "sim", "s", "contribuinte"):
        return True
    if s in ("0", "false", "não", "nao", "n", "9", "2"):
        return False
    return None


def ingest_sheet(data, filename, mapping: dict, sheet=None, header_rows=1):
    if not {"issued", "description", "value"} <= mapping.keys():
        raise ValueError("Mapeie pelo menos data, descrição e valor do produto")
    if set(mapping) - FIELDS.keys() or any(
        not isinstance(i, int) or i < 0 for i in mapping.values()
    ):
        raise ValueError("Mapeamento inválido")
    books, warnings, partial = workbook(data, filename)
    sheet = sheet or next(iter(books))
    if sheet not in books:
        raise ValueError("Aba não encontrada")
    result = Ingestion(warnings=warnings, partial=partial)
    grouped = {}
    for line, row in enumerate(books[sheet][header_rows:], header_rows + 1):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        try:
            values = {field: row[col] if col < len(row) else None for field, col in mapping.items()}
            raw_date = values["issued"]
            if isinstance(raw_date, datetime):
                issued = raw_date.date()
            elif isinstance(raw_date, date):
                issued = raw_date
            else:
                try:
                    issued = date.fromisoformat(str(raw_date)[:10])
                except ValueError:
                    issued = datetime.strptime(str(raw_date), "%d/%m/%Y").date()
            key, number, series = (string(values.get(k)) for k in ("key", "number", "series"))
            group = key or (f"{issued}/{series}/{number}" if number else f"linha:{line}")
            fields = {}
            for field in ("code", "description", "ncm", "cest", "cfop", "unit"):
                fields[field] = string(values.get(field))
            fields["description"] = fields["description"] or ""
            for field in (
                "quantity",
                "unit_value",
                "value",
                "discount",
                "freight",
                "insurance",
                "other",
            ):
                fields[field] = decimal(values.get(field))
            for field, val in values.items():
                if "." in field:
                    t, f = field.split(".")
                    v = string(val) if f in ("cst", "csosn") else decimal(val)
                    if v and f == "cst":
                        v = v.zfill(2)
                    fields.setdefault(t, {})[f] = v
            n = (
                int(values["n_item"])
                if values.get("n_item")
                else len(grouped.get(group, {}).get("items", [])) + 1
            )
            item = Item(**fields, n_item=n, source={"file": filename, "sheet": sheet, "line": line})
            head = dict(
                key=key,
                number=number,
                series=series,
                issued=issued,
                model=int(values["model"]) if values.get("model") else None,
                recipient_taxpayer=boolean(values.get("recipient_taxpayer")),
                final_consumer=boolean(values.get("final_consumer")),
                recipient_uf=string(values.get("recipient_uf")),
                operation=string(values.get("operation")) or "desconhecida",
                status=string(values.get("status")) or "desconhecida",
            )
            if group not in grouped:
                grouped[group] = {**head, "items": []}
            elif any(grouped[group][k] != v for k, v in head.items()):
                raise ValueError("Dados de cabeçalho conflitantes para a mesma nota")
            if n in [i.n_item for i in grouped[group]["items"]]:
                raise ValueError("Item duplicado na mesma nota")
            grouped[group]["items"].append(item)
            result.read += 1
        except (ValueError, TypeError, InvalidOperation, OverflowError) as exc:
            result.discarded += 1
            result.partial = True
            result.errors.append(
                {"file": filename, "sheet": sheet, "line": line, "reason": str(exc)[:300]}
            )
    result.documents = [Document(**d) for d in grouped.values()]
    if not result.documents:
        result.partial = True
    if any(d.status == "desconhecida" or d.operation == "desconhecida" for d in result.documents):
        result.partial = True
        result.warnings.append(
            "Situação/operação não informada em parte das linhas; regras fiscais limitadas."
        )
    result.warnings.append(
        "Cobertura limitada ao arquivo; datas ausentes não comprovam dias sem movimento."
    )
    return result
