from decimal import Decimal

from auditoria.confirmed_report import publish
from auditoria.engine import evaluate


def report(document, profile):
    return publish(
        evaluate([document], [profile], []),
        {
            "documents": [document.model_dump(mode="json")],
            "profiles": [profile.model_dump(mode="json")],
        },
    )


def test_publication_removes_guesses_without_mutating_candidates(document, profile):
    document.items[0].cfop = "5927"
    candidates = evaluate([document], [profile], [])
    assert any(f.code == "C08" for f in candidates.findings)
    output = publish(candidates)
    assert not output.findings
    assert candidates.findings
    assert publish(output) == output


def test_same_line_two_taxes_counts_once(document, profile):
    profile.metodo_pis_cofins = "com_exclusao_icms"
    result = report(document, profile)
    base = result.summary["base_review"]
    assert len(result.findings) == 2
    assert result.summary["findings"] == 1
    assert (base["missing_icms"], base["other_difference"], base["compatible"]) == (1, 0, 0)
    assert base["rows"][0]["expected_base"] == "82.00"
    assert publish(result) == result


def test_partial_missing_base_never_claims_compatible(document, profile):
    profile.metodo_pis_cofins = "com_exclusao_icms"
    document.items[0].pis.base = Decimal("82")
    document.items[0].cofins.base = None
    base = report(document, profile).summary["base_review"]
    assert base["compatible"] == 0 and base["unassessed"] == 1


def test_compatible_and_other_composition(document, profile):
    profile.metodo_pis_cofins = "com_exclusao_icms"
    document.items[0].pis.base = document.items[0].cofins.base = Decimal("82")
    assert report(document, profile).summary["base_review"]["compatible"] == 1
    document.items[0].other = Decimal("3.20")
    base = report(document, profile).summary["base_review"]
    assert base["other_difference"] == 1
    assert base["rows"][0]["expected_base"] == "85.20"


def test_missing_method_components_and_negative_base_do_not_pass(document, profile):
    assert report(document, profile).summary["base_review"]["unassessed"] == 1
    profile.metodo_pis_cofins = "com_exclusao_icms"
    document.items[0].discount = None
    assert report(document, profile).summary["base_review"]["unassessed"] == 1
    document.items[0].discount = Decimal("99")
    assert report(document, profile).summary["base_review"]["unassessed"] == 1


def test_cancelled_excluded_and_single_eligible_tax(document, profile):
    profile.metodo_pis_cofins = "com_exclusao_icms"
    document.items[0].cofins.cst = "49"
    assert report(document, profile).summary["base_review"]["missing_icms"] == 1
    document.status = "cancelada"
    base = report(document, profile).summary["base_review"]
    assert base["missing_icms"] == base["compatible"] == base["unassessed"] == 0


def test_cst_contradiction_is_retained(document, profile):
    document.items[0].icms.cst = "41"
    output = report(document, profile)
    assert [f.code for f in output.findings] == ["C03"]


def test_missing_other_tax_base_keeps_known_difference(document, profile):
    profile.metodo_pis_cofins = "com_exclusao_icms"
    document.items[0].cofins.base = None
    base = report(document, profile).summary["base_review"]
    assert base["missing_icms"] == 1
    assert base["rows"][0]["cofins_base"] is None


def test_concise_downloads_have_summary_and_one_row_per_item(document, profile):
    from io import BytesIO

    from openpyxl import load_workbook

    from auditoria.reports import html_report, xlsx_report

    profile.metodo_pis_cofins = "com_exclusao_icms"
    result = report(document, profile)
    html = html_report(result, {})
    from lxml.html import fromstring

    text = " ".join(fromstring(html).itertext())
    assert "1 linha sem descontar o ICMS da base" in text
    assert "R$ 82,00" in html
    book = load_workbook(BytesIO(xlsx_report(result, {})))
    assert book["Bases PIS COFINS"].max_row == 2
    assert book["Ocorrências"].max_row == 1
    assert "Não avaliadas" not in book.sheetnames
