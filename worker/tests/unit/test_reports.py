from io import BytesIO

import openpyxl

from auditoria.engine import evaluate
from auditoria.reports import html_report, pdf_report, xlsx_report


def test_xlsx_formula_injection_is_escaped(document, profile):
    document.items[0].description = '=HYPERLINK("https://evil.invalid")'
    document.items[0].ncm = "00000000"
    document.items[0].icms.cst = "41"
    result = evaluate([document], [profile], [])
    data = xlsx_report(result, {"Cliente": "=1+1"})
    book = openpyxl.load_workbook(BytesIO(data))
    assert book["Resumo"]["B1"].value.startswith("'")
    assert book["Ocorrências"]["D2"].data_type == "s"
    assert book["Ocorrências"]["D2"].value.startswith("'")
    assert "Não avaliadas" not in book.sheetnames


def test_pdf_and_html_escape_customer_text(document, profile):
    result = evaluate([document], [profile], [])
    context = {"Cliente": "<script>alert(1)</script>", "Cobertura": "Parcial"}
    html = html_report(result, context)
    assert "<script>" not in html and "&lt;script&gt;" in html
    pdf = pdf_report(result, context)
    assert pdf.startswith(b"%PDF") and len(pdf) > 2000


def test_pdf_with_embedded_logo(document, profile):
    import base64

    from PIL import Image

    from auditoria.api import BrandingInput

    image = BytesIO()
    Image.new("RGB", (10, 10), "#173d35").save(image, format="PNG")
    logo = "data:image/png;base64," + base64.b64encode(image.getvalue()).decode()
    branding = BrandingInput(nome="Escritório", cor="#173d35", logo_data_url=logo)
    pdf = pdf_report(
        evaluate([document], [profile], []), {"Cliente": "Teste"}, branding.model_dump()
    )
    assert pdf.startswith(b"%PDF")


def test_branding_rejects_external_resource():
    import pytest

    from auditoria.api import BrandingInput

    with pytest.raises(ValueError):
        BrandingInput(nome="Teste", logo_data_url="file:///etc/passwd")


def test_pdf_product_sections_show_findings(document, profile):
    document.items[0].ncm = "00000000"
    document.items[0].icms.cst = "41"
    result = evaluate([document], [profile], [])
    html = html_report(result, {"Cliente": "Teste"})
    assert "NCM coringa" not in html and "Produto sintético" in html
    assert "Outras divergências documentais" in html
    assert pdf_report(result, {"Cliente": "Teste"}).startswith(b"%PDF")


def test_operation_report_keeps_non_error_rows_and_missing_values(document, profile):
    document.items[0].value = None
    document.items[0].cfop = "5202"
    result = evaluate([document], [profile], [])
    book = openpyxl.load_workbook(BytesIO(xlsx_report(result, {})))
    assert book["Itens conferidos"].max_row == 2
    assert book["Itens conferidos"]["M2"].value is None
    assert book["Por CFOP e categoria"]["B2"].value.startswith("Devolução")
    assert "Operação de natureza específica" not in html_report(result, {})


def test_client_report_preserves_every_row_across_table_sections(document, profile):
    from lxml.html import fromstring

    from auditoria.confirmed_report import publish

    profile.metodo_pis_cofins = "com_exclusao_icms"
    original = document.items[0]
    document.items = [
        original.model_copy(update={"n_item": i, "source": {"line": 100 + i}}) for i in range(1, 26)
    ]
    result = publish(
        evaluate([document], [profile], []),
        {
            "documents": [document.model_dump(mode="json")],
            "profiles": [profile.model_dump(mode="json")],
        },
    )
    tree = fromstring(html_report(result, {"Cliente": "Cliente sintético"}))
    rows = tree.xpath("//tbody/tr")
    assert len(rows) == 25
    assert [row.xpath("./td[1]")[0].text_content().split("Linha ")[1] for row in rows] == [
        str(100 + i) for i in range(1, 26)
    ]
    assert "Encaminhamento ao responsável fiscal" in tree.text_content()
    assert "25" in tree.text_content()
